"""Geometric QA for the EKADHARA deck.

Substitute for image-based visual QA (no LibreOffice available on this host).
Checks: out-of-bounds shapes, edge margins, text-box overlaps, estimated text overflow.
"""
import sys
from pptx import Presentation
from pptx.util import Emu

EMU_IN = 914400.0
SLIDE_W, SLIDE_H = 13.333, 7.5
MARGIN = 0.5

# rough advance-width ratio (fraction of font size) per family
WIDTH_RATIO = {"Calibri": 0.47, "Cambria": 0.50, "Courier New": 0.60}


def inches(v):
    return (v or 0) / EMU_IN


def box(sh):
    return (inches(sh.left), inches(sh.top),
            inches(sh.left) + inches(sh.width),
            inches(sh.top) + inches(sh.height))


def overlap(a, b):
    ix = min(a[2], b[2]) - max(a[0], b[0])
    iy = min(a[3], b[3]) - max(a[1], b[1])
    return ix * iy if (ix > 0 and iy > 0) else 0.0


def est_overflow(sh):
    """Estimate whether text needs more vertical space than the box provides."""
    if not sh.has_text_frame:
        return None
    tf = sh.text_frame
    text = tf.text
    if not text.strip():
        return None
    w_in = inches(sh.width)
    h_in = inches(sh.height)
    if w_in <= 0.05 or h_in <= 0.05:
        return None

    total_h = 0.0
    for para in tf.paragraphs:
        ptext = "".join(r.text for r in para.runs)
        sizes = [r.font.size.pt for r in para.runs if r.font.size]
        fam = next((r.font.name for r in para.runs if r.font.name), "Calibri")
        fs = max(sizes) if sizes else 12.0
        ratio = WIDTH_RATIO.get(fam, 0.47)
        cpl = max(1, int((w_in * 72.0) / (fs * ratio)))
        # honour explicit newlines inside the paragraph
        segs = ptext.split("\n") if ptext else [""]
        lines = sum(max(1, -(-len(s) // cpl)) for s in segs)
        total_h += lines * fs * 1.22 / 72.0
    return total_h, h_in


def main(path):
    prs = Presentation(path)
    problems = 0

    for idx, slide in enumerate(prs.slides, start=1):
        print(f"\n=== SLIDE {idx} ===")
        text_boxes = []

        for sh in slide.shapes:
            if sh.left is None:
                continue
            l, t, r, b = box(sh)
            label = (sh.text_frame.text[:42].replace("\n", " ")
                     if sh.has_text_frame and sh.text_frame.text.strip()
                     else f"<{sh.shape_type}>")

            # bounds
            if r > SLIDE_W + 0.01 or b > SLIDE_H + 0.01 or l < -0.01 or t < -0.01:
                print(f"  OUT-OF-BOUNDS  ({l:.2f},{t:.2f})-({r:.2f},{b:.2f})  {label!r}")
                problems += 1
            # edge margin
            elif l < MARGIN - 0.01 or t < MARGIN - 0.01 or \
                    r > SLIDE_W - MARGIN + 0.01 or b > SLIDE_H - MARGIN + 0.01:
                # footers legitimately sit low; only warn above the footer band
                if b < SLIDE_H - 0.55:
                    print(f"  tight margin   ({l:.2f},{t:.2f})-({r:.2f},{b:.2f})  {label!r}")

            # overflow
            ov = est_overflow(sh)
            if ov:
                need, have = ov
                if need > have * 1.06:
                    print(f"  TEXT OVERFLOW  needs {need:.2f}\" has {have:.2f}\"  {label!r}")
                    problems += 1

            if sh.has_text_frame and sh.text_frame.text.strip():
                text_boxes.append((box(sh), label, inches(sh.width) * inches(sh.height)))

        # text-on-text overlap (ignore small incidental overlaps)
        for i in range(len(text_boxes)):
            for j in range(i + 1, len(text_boxes)):
                a, la, aa = text_boxes[i]
                bb, lb, ab = text_boxes[j]
                ov = overlap(a, bb)
                if ov > 0.06 * min(aa, ab) and ov > 0.05:
                    print(f"  TEXT OVERLAP   {ov:.2f} sq in :: {la!r} <> {lb!r}")
                    problems += 1

    print(f"\n{'='*60}\nPROBLEMS: {problems}")
    return 0 if problems == 0 else 1


if __name__ == "__main__":
    sys.exit(main(sys.argv[1]))
